import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

# Create a new presentation
prs = Presentation()

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Bankism – Blue‑Team Defensive AI Firewall"
subtitle.text = "Project Overview & Architecture"

# Slide 2: Architecture Diagram (placeholder image)
slide = prs.slides.add_slide(prs.slide_layouts[5])
title = slide.shapes.title
title.text = "Architecture Diagram"
# Add placeholder for diagram
left = Inches(1)
top = Inches(2)
width = Inches(8)
height = Inches(4.5)
# Assuming diagram saved as architecture.png in repo root (we'll create a simple placeholder later)
img_path = Path('architecture.png')
if img_path.exists():
    slide.shapes.add_picture(str(img_path), left, top, width=width, height=height)
else:
    # Add a textbox indicating missing image
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = "[Architecture diagram will be placed here]"
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18)

# Slide 3: Request Flow & Workflow
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Request Flow & Workflow"
content = slide.placeholders[1].text_frame
content.text = "1️⃣ User types a request in the UI\n2️⃣ Debounced → POST /api/preflight\n3️⃣ BankismFirewallPipeline (normaliser, PII redactor, semantic guard)\n4️⃣ JSON verdict returned\n5️⃣ If approved → BankingAgent → Bank Engine\n6️⃣ Updates UI, metrics, events\n7️⃣ Blocked requests are logged to SOC"

# Slide 4: UI Features
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Key UI Features"
content = slide.placeholders[1].text_frame
content.text = "• Sign‑in overlay with demo accounts\n• Currency selector (USD, EUR, GBP, INR, JPY, AUD, CAD)\n• Live reactive dashboard (140 ms debounce)\n• AI Guide‑Bot (Gemini fallback)\n• MFA limit control and event inspector"

# Slide 5: Getting Started & Deployment
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
title.text = "Getting Started & Deployment"
content = slide.placeholders[1].text_frame
content.text = "1️⃣ Install dependencies: pip install -r requirements.txt\n2️⃣ Run server: python run_server.py (auto‑opens Chrome)\n3️⃣ Access dashboard at http://127.0.0.1:8000\n4️⃣ Use demo accounts (alice/alice123, bob/bob123, mallory/mallory123)\n5️⃣ Extend firewall pipeline for production use"

# Save presentation
output_path = Path('Bankism_Overview.pptx')
prs.save(str(output_path))
print(f"Presentation saved to {output_path}")
